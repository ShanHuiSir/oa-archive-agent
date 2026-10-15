"""Generate original, clearly labeled fictitious office documents."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from openpyxl import Workbook
from PIL import Image,ImageDraw,ImageFont
import fitz
ROOT=Path(__file__).resolve().parents[1];S=ROOT/'samples';S.mkdir(exist_ok=True)
def word(name,title,lines):
    d=Document();st=d.styles['Normal'];st.font.name='Microsoft YaHei';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑');st.font.size=Pt(12)
    d.add_heading(title,0)
    for line in ['模拟业务数据 仅用于软件功能演示',*lines]:d.add_paragraph(line)
    d.save(S/name)
word('01_值班通知.docx','节假日值班安排通知',['发文部门：行政部','签发日期：2026-04-20','为做好假期服务工作，综合服务岗按排班表值班。联系人及人员均为虚构。','文件类型：通知'])
w=Workbook();s=w.active;s.title='模拟预算';
for r in [['模拟业务数据 仅用于软件功能演示'],['季度预算执行报表'],['所属部门','财务部'],['生成日期','2026-03-31'],['项目','预算元','执行元'],['办公用品',12000,9600],['培训',8000,6400],['合计','=SUM(B6:B7)','=SUM(C6:C7)']]:s.append(r)
s.column_dimensions['A'].width=38;s.column_dimensions['B'].width=18;s.column_dimensions['C'].width=18;w.save(S/'02_预算执行.xlsx')
d=fitz.open();p=d.new_page();p.insert_text((60,75),'模拟业务数据 仅用于软件功能演示',fontname='china-s',fontsize=15)
for i,t in enumerate(['办公设备采购合同','归档部门：法务部','签署日期：2026-02-18','甲方：星禾模拟企业；乙方：明川模拟设备商。','标的：演示用办公设备。合同金额人民币24000元。','本文件不构成真实交易，企业名称与业务事项均为虚构。']):p.insert_text((60,125+40*i),t,fontname='china-s',fontsize=14)
d.save(S/'03_设备合同.pdf');d.close()
im=Image.new('RGB',(1500,1100),'white');draw=ImageDraw.Draw(im);font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',38)
for i,t in enumerate(['模拟业务数据 仅用于软件功能演示','研发部项目推进会议纪要','所属部门：研发部','会议纪要生成日期：2026-03-12','会议结论：先完成设备数据接入，再核对异常指标。','参会人：虚构人员甲、虚构人员乙。']):draw.text((90,90+i*120),t,fill='#162f36',font=font)
im.save(S/'04_扫描会议纪要.png')
word('05_企业自定义模板.docx','安全检查报告',['模拟企业自定义 OA 模板','归口单位：生产部','制发日期：2026-04-08','材料类别：报告','检查事项：消防通道、设备点检、劳保用品。','检查结果：模拟检查表记录两项整改事项，已由模拟责任人复核。'])
(S/'06_日期缺失.txt').write_text('模拟业务数据 仅用于软件功能演示\n销售部工作总结报告\n所属部门：销售部\n文件类型：报告\n完成渠道资料整理及模拟业务复盘。本文没有生成日期，归档前应由资料负责人核对。\n',encoding='utf-8')
expected=[dict(file=n,department=dep,type=t,date=dt) for n,dep,t,dt in [
('01_值班通知.docx','行政部','通知','2026-04-20'),('02_预算执行.xlsx','财务部','报表','2026-03-31'),('03_设备合同.pdf','法务部','合同','2026-02-18'),('04_扫描会议纪要.png','研发部','会议纪要','2026-03-12'),('05_企业自定义模板.docx','生产部','报告','2026-04-08'),('06_日期缺失.txt','销售部','报告','')]]
(ROOT/'tests'/'expected_samples.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2),encoding='utf-8')
print('Generated',len(expected),'fictitious samples')
