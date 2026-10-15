"""Check actual screen-run outputs independently of app's status messages."""
from pathlib import Path
import json,zipfile,hashlib,sys
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from oa.core import digest
r=json.loads((ROOT/'demo/本轮操作结果.json').read_text(encoding='utf-8'));rows=r['ledger'];items=r['state']['items'];assert len(rows)==len(items)==6
sources={p.name:p for p in (ROOT/'samples').iterdir() if p.is_file()};checks=[]
with zipfile.ZipFile(ROOT/'demo/本轮归档成果.zip') as z:
    manifests=json.loads(z.read('归档记录.json'));assert len(manifests)==6
    for row in rows:
        data=z.read('归档文件/'+row['destination'].replace('\\','/'))
        actual=hashlib.sha256(data).hexdigest();expected=digest(sources[row['source_name']]);assert actual==row['sha256']==expected
        checks.append(dict(file=row['source_name'],sha256=actual,copy_equals_original=True))
    names=z.namelist();assert not any(n.startswith(('/', '..')) for n in names)
for name,count in [('本轮文件目录台账.xlsx',6),('筛选台账_法务部.xlsx',1)]:
    w=load_workbook(ROOT/'demo'/name);assert w.active.max_row-1==count
    if count==1:assert w.active['A2'].value=='法务部'
    w.close()
assert next(x for x in items if x['source']=='06_日期缺失.txt')['date']=='2026-04-09'
assert all(x['model']=='local-rules-v1' and not x['request_id'] for x in items)
report=dict(version='1.0.0',input_count=6,archived_count=6,ledger_count=6,filtered_count=1,all_copy_hashes_match=True,checks=checks)
(ROOT/'validation/demo_integrity.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print('Verified six actual archives and both exported ledgers')
