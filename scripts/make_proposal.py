"""Build both private and redacted editable proposals from measured results."""
from pathlib import Path
import json,argparse,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]

def font(run,size=11,bold=False,name='宋体'):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
def para(d,text='',size=11):
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.35;font(p.add_run(text),size);return p
def heading(d,text):
    p=d.add_paragraph(text,style='Heading 1');p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(10)
    for r in p.runs:font(r,15,True,'黑体')
    return p
def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]);tc=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EEEC' if ri==0 else 'FFFFFF');tc.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5);p.paragraph_format.line_spacing=1.15
                for r in p.runs:font(r,10,ri==0)
    para(d,'',3);return t

def build(path,info):
    live=json.loads((ROOT/'validation/local_samples.json').read_text(encoding='utf-8'))
    run=json.loads((ROOT/'demo/本轮操作结果.json').read_text(encoding='utf-8'))
    vid=json.loads((ROOT/'demo/视频实测记录.json').read_text(encoding='utf-8'))
    att=json.loads((ROOT/'validation/attachment_summary.json').read_text(encoding='utf-8'))
    tests=json.loads((ROOT/'validation/test_summary.json').read_text(encoding='utf-8'))
    rows=run['ledger'];items=run['state']['items']
    checks=sum(sum(x.get('field_checks',{}).values()) for x in live['records']);check_count=sum(len(x.get('field_checks',{})) for x in live['records'])
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='宋体';st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        st.font.italic=False;st.font.underline=False
    for el in list(d.styles._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
    d.core_properties.title='企业OA文件智能归档智能体项目申报书';d.core_properties.author='项目申报人';d.core_properties.version='1.0.0'
    p=d.add_paragraph(style='Title');font(p.add_run('企业OA文件智能归档智能体'),22,True,'黑体');p.paragraph_format.space_after=Pt(8)
    p=d.add_paragraph(style='Subtitle');font(p.add_run('归序 项目申报书'),16);p.paragraph_format.space_after=Pt(18)
    para(d,'作品版本  1.0.0    申报日期  2026年10月15日')
    table(d,['参赛信息','填写内容'],[['姓名',info['name']],['单位',info['organization']],['联系方式',info['contact']]], [3.2,13.4])
    heading(d,'项目简介')
    para(d,'归序面向企业 OA 文件整理工作，把“导入资料、识别信息、核对归档、查找台账、导出成果”放在一个本地图形界面中。用户不必逐个建立文件夹、修改文件名，也能够看到每份文件的识别依据和最终存放位置。')
    para(d,'我选择这个题目，是因为办公资料经常混在同一个文件夹里：通知、合同、表格和扫描件的命名习惯不一致，过一段时间就很难查找。我希望把整理过程做得明确一些，让使用者既能批量处理，也能在有疑点时停下来核对。')
    heading(d,'已完成成果')
    para(d,f'项目已完成可运行应用、Windows 独立程序、完整源码及使用说明、原创模拟样例、测试记录、中文配音与同步字幕演示视频。本轮演示导入 {len(items)} 份模拟办公资料，经过本地规则识别、真实 OCR 与人工复核，最终归档 {len(rows)} 份，导出 Excel 文件目录台账和包含文件副本的归档 ZIP。')
    para(d,'项目中的演示企业、人员和业务事项均为模拟。比赛附件中的十份文件作为本机输入验证资料，未作为真实企业生产数据，也未用示例文本代替程序处理结果。')
    d.add_page_break();heading(d,'一 用户操作与业务规则')
    table(d,['步骤','用户操作','实际功能'],[
        ['导入资料','选择本地文件夹或多份文件','读取 Word、Excel、PDF、图片和 TXT，显示输入列表'],
        ['识别信息','点击本地智能识别或模型增强识别','提取正文及 OCR；本地规则定位字段，模型增强则真实请求配置的接口'],
        ['核对疑点','查看正文、原文引文和字段','缺失日期、依据不完整或低置信结果进入人工复核'],
        ['分类归档','指定输出路径并确认','复制原件，按规则生成目录和文件名，核验内容哈希'],
        ['查询导出','搜索关键词并导出','生成筛选后的 Excel 台账或全量归档包']
    ],[2.1,5.2,9.3])
    heading(d,'按企业模板识别')
    para(d,'设置页允许维护部门、文件类型、命名模板和 OA 字段含义。演示中的自定义模板使用“归口单位”“制发日期”“材料类别”，配置后分别解释为部门、生成日期和文件类型。命名规则支持部门、类型、日期、标题四个占位符。')
    para(d,'默认目录为“部门／文件类型／年月”，文件名由部门、类型、生成日期和标题组成。分类字段先按配置枚举核验；日期必须是真实的年月日格式。程序不会把当前日期或合同到期日直接当成生成日期。')
    heading(d,'人工复核与原件保护')
    para(d,'模型自评置信度只用于提醒核对，不作为准确率。缺失日期样例在识别后保持日期为空；演示由模拟资料负责人补录 2026年4月9日，再确认归档。这个日期是人工提供的信息，不计作模型自动识别结果。')
    para(d,'归档采用复制方式，保留输入原件。相同内容按 SHA256 跳过重复，同名不同内容增加后缀。识别后原件若被修改，程序拒绝直接归档，避免识别信息与文件内容错配。')
    d.add_page_break();heading(d,'二 技术架构与工作流')
    para(d,'项目采用 Python 实现本地服务，浏览器提供图形工作台。Windows 独立程序将运行环境、OCR 模型和应用代码一并打包，启动后打开回环地址。识别在后台线程执行，用户可以持续查看进度和处理日志。本地规则模式无需 API Key；模型增强模式使用用户本机配置。')
    table(d,['组成','已经实现的职责'],[
        ['文件提取工具','DOCX 段落与表格、Excel 工作表、PDF 正文；图片和扫描页通过本地 RapidOCR 提取'],
        ['规则及模型客户端','本地规则提取部门、类型与日期；可选模型客户端将正文和配置送入 DeepSeek，校验结构化响应并处理超时或服务错误'],
        ['核验与复核','检查字段枚举、日期、引文和置信阈值；支持在界面查看正文并修正'],
        ['归档与台账','本地规则构造路径，复制与哈希校验，SQLite 记录元数据，导出 Excel 和 ZIP']
    ],[3.5,13.1])
    para(d,'Agent 工作流固定为“发现文件 → 提取正文 → 规则识别或模型增强 → 本地核验 → 必要时人工复核 → 规则归档 → 台账导出”。规则候选与模型结果分别记录来源。模型只能给出候选字段，不能执行文件内的命令，也不能自行决定物理输出路径。')
    heading(d,'数据处理边界')
    para(d,'OCR 和归档在本机进行；仅选择模型增强时，将分类所需文件名与提取文字发送到用户配置的模型服务，二进制原件不上传。API Key 从本机配置或环境变量读取，不回显到设置页，不进入台账、视频和公开交付包。服务仅监听 127.0.0.1，写入接口校验会话令牌及请求来源。')
    para(d,'程序限制单文件不超过 80 MB。PDF 最多读取前12页，Excel 每表读取前1000行，模型正文最多18000字符；触及截取范围会提示复核，归档副本仍保留完整文件。损坏、加密或无法提取的输入显示失败提示。')
    d.add_page_break();heading(d,'三 本轮验证结果')
    para(d,'验证于2026年10月15日在 Windows、Python 3.12 环境完成。下表来自实际程序运行记录；原创样例的预期字段在处理前定义，不能推广为所有企业文档的分类准确率。')
    table(d,['输入样例','格式','部门','类型','生成日期'],[[x['source'].split('_',1)[1],Path(x['source']).suffix,x['department'],x['type'],x['date'] or '空值 待复核'] for x in live['records']],[4.5,1.7,2.4,2.3,5.7])
    para(d,f'原创模拟样例共6份，实际使用本地规则模式，18个预期字段中本轮符合 {checks} 项。空日期的预期是保持空值并进入复核。提取方式、处理耗时与逐字段检查保存在 validation/local_samples.json。这是对固定样例的核对结果，不是模型准确率。')
    para(d,f'另对比赛附件的 {att["input_count"]} 份办公文件执行本地提取、规则识别和归档检查。附件没有官方标注结果，因此保留运行状态和提取方式，不据此宣称官方准确率。完整输入及详情保留在本机，公开版本只保留不含正文的汇总。')
    para(d,'模型客户端使用提供的有效账号真实请求官方接口，本轮服务返回503或持续保活后超时。接口联调记录与本地规则结果分开保留；本轮演示、字段核对和归档成果均由本地规则模式产生，不作为成功模型响应的证明。')
    table(d,['核验项目','本轮结果'],[
        ['本地自动化回归',f'{tests["tests_run"]} 项通过，覆盖去重、冲突、原件变化、复核、路径与导出保护'],
        ['演示归档与导出',f'{len(rows)} 份归档文件；全量台账 {len(rows)} 条；法务部筛选台账1条'],
        ['文件完整性','对本轮归档副本逐份比对 SHA256，与输入原件一致'],
        ['视频验证',f'MP4 时长 {vid["duration_s"]:.1f} 秒，中文神经配音与同步字幕，实际浏览器操作录制']
    ],[4.5,12.1])
    para(d,'本轮记录支持上述格式和案例的功能闭环结论，不据此推算用户规模、长期稳定性或效率提升比例。全部数量均以本次保留的运行记录为依据。')
    d.add_page_break();heading(d,'四 实际界面与交付使用')
    image=ROOT/'docs/台账实际结果.png'
    if image.exists():d.add_picture(str(image),width=Cm(16.5));para(d,'图1 本轮模拟案例真实归档后的文件目录台账',10)
    heading(d,'安装与使用')
    para(d,'解压完整源码 ZIP，进入 Windows／OAArchive 目录运行 OAArchive.exe。本地模式无需 API Key，导入 samples 中的模拟资料，点击本地智能识别，核对后归档，再进入台账搜索和导出。需要模型增强时，在设置页填写本机 API Key 与可用模型名。关闭服务控制台即可结束程序。')
    para(d,'使用源码时，按 README 创建 Python 虚拟环境并安装依赖，再运行 main.py。项目提供安装、启动、打包入口和自动化测试。架构说明中列出处理范围、工作流和数据流向，便于复现与检查。')
    heading(d,'交付文件')
    para(d,'完整源码 ZIP 包含工程、Windows 可执行程序、依赖与配置示例、使用说明、原创输入样例、必要测试、实际运行记录，以及解说词和字幕。演示 MP4 与可编辑申报书单独提供；公开申报书与 ZIP 隐藏参赛人的私人信息，本地提交版保留完整参赛信息。')
    para(d,'我希望这件作品首先把一个具体问题做好：让使用者能够批量整理文件，也能解释每份文件为什么这样归类。归档结果能在目录中直接查找，识别依据和人工复核记录也能随台账继续保留。')
    path.parent.mkdir(parents=True,exist_ok=True);d.save(path);print(path)

def main():
    p=argparse.ArgumentParser();p.add_argument('--private-info',required=True);args=p.parse_args()
    build(ROOT/'docs/项目申报书公开版.docx',dict(name='公开版已隐藏',organization='公开版已隐藏',contact='公开版已隐藏'))
    build(ROOT/'delivery/项目申报书.docx',json.loads(Path(args.private_info).read_text(encoding='utf-8')))
if __name__=='__main__':main()
