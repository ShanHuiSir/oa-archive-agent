import json,tempfile,unittest,sys
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from oa.core import *

class ArchiveTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.base=Path(self.tmp.name);self.c=DEFAULT.copy();self.source=self.base/'来源.txt';self.source.write_text('行政部 通知 2026-04-20',encoding='utf-8')
    def tearDown(self):self.tmp.cleanup()
    def item(self,title='测试通知'):
        return Item(str(self.source),sha256=digest(self.source),department='行政部',type='通知',date='2026-04-20',title=title,status='可归档')
    def test_lossless_duplicate_search_export(self):
        ledger=Ledger(self.base/'archive');x=self.item();r=ledger.archive(x,self.c)
        self.assertEqual(digest(self.source),digest(ledger.root/r['destination']));self.assertIsNone(ledger.archive(self.item(),self.c))
        self.assertEqual(len(ledger.rows('通知')),1);self.assertEqual(len(ledger.rows('不存在')),0)
        out=self.base/'台账.xlsx';self.assertEqual(ledger.export(out),1)
        from openpyxl import load_workbook
        w=load_workbook(out);self.assertEqual(w.active.max_row,2);self.assertEqual(w.active['G2'].value,digest(self.source));w.close()
    def test_source_mutation_blocked(self):
        x=self.item();self.source.write_text('changed');
        with self.assertRaisesRegex(ValueError,'发生变化'):Ledger(self.base/'a').archive(x,self.c)
        self.assertEqual(Ledger(self.base/'a').rows(),[])
    def test_review_required(self):
        x=self.item();x.status='待复核'
        with self.assertRaises(ValueError):Ledger(self.base/'a').archive(x,self.c)
    def test_traversal_cannot_escape(self):
        x=self.item('../../CON');p=target_relative(x,self.c)
        self.assertTrue((self.base/p).resolve().is_relative_to(self.base));self.assertNotIn('..',p.parts)
    def test_collision_preserves_both(self):
        ledger=Ledger(self.base/'a');r=ledger.archive(self.item(),self.c)
        self.source.write_text('another',encoding='utf-8');r2=ledger.archive(self.item(),self.c)
        self.assertNotEqual(r['destination'],r2['destination']);self.assertEqual(len(ledger.rows()),2)
    def test_formula_injection(self):
        self.assertEqual(excel_safe('=HYPERLINK("x")')[0],"'");self.assertEqual(excel_safe('normal'),'normal')
    def test_invalid_template(self):
        for name in ['{source}','{department.__class__}']:
            c=self.c.copy();c['name_template']=name
            with self.assertRaises(ValueError):validate_config(c)
    def test_no_key_gives_failure(self):
        with patch('oa.core.classify',side_effect=ValueError('请设置 API Key')):x=analyze(self.source,self.c)
        self.assertEqual(x.status,'失败');self.assertNotEqual(x.department,'行政部')
    def test_missing_date_requires_review(self):
        raw=dict(department='行政部',type='通知',date='',title='测试',confidence=.5,evidence=dict(department='行政部',type='通知',date=''))
        with patch('oa.core.classify',return_value=(raw,dict(model='test',request_id='test',usage={}))):x=analyze(self.source,self.c)
        self.assertEqual(x.status,'待复核');self.assertEqual(x.date,'')
    def test_discover_excludes_archive_and_unsupported(self):
        out=self.base/'a';out.mkdir();(out/'x.txt').write_text('x');(self.base/'x.exe').write_text('x')
        self.assertEqual(discover(self.base,out),[self.source])
    def test_corrupt_file_reports_failure(self):
        p=self.base/'bad.pdf';p.write_bytes(b'not pdf');x=analyze(p,self.c);self.assertEqual(x.status,'失败');self.assertTrue(x.notes)
    def test_local_missing_date_is_empty(self):
        self.source.write_text('行政部\n文件类型：通知\n无日期资料',encoding='utf-8');x=analyze_local(self.source,self.c)
        self.assertEqual(x.date,'');self.assertEqual(x.status,'待复核');self.assertEqual(x.model,'local-rules-v1');self.assertEqual(x.request_id,'')
    def test_local_conflicting_dates_review(self):
        self.source.write_text('行政部 通知\n生成日期：2026-04-20\n签发日期：2026-04-21',encoding='utf-8');x=analyze_local(self.source,self.c)
        self.assertEqual(x.status,'待复核');self.assertEqual(x.date,'')
    def test_local_custom_template(self):
        self.source.write_text('安全检查报告\n归口单位：生产部\n制发日期：2026年4月8日\n材料类别：报告',encoding='utf-8');x=analyze_local(self.source,self.c)
        self.assertEqual((x.department,x.type,x.date,x.status),('生产部','报告','2026-04-08','可归档'))
    def test_local_filename_type_over_body_reference(self):
        p=self.base/'财务部_预算汇总表_2026-04-20.txt';p.write_text('为落实管理制度，特编制本统计资料。',encoding='utf-8');x=analyze_local(p,self.c)
        self.assertEqual(x.type,'报表')
    def test_local_configurable_aliases(self):
        self.c['template_hint']='承办科室代表部门；落款时间代表生成日期；资料类别代表文件类型。'
        self.source.write_text('检查材料\n承办科室：生产部\n落款时间：2026年4月8日\n资料类别：报告',encoding='utf-8');x=analyze_local(self.source,self.c)
        self.assertEqual((x.department,x.type,x.date,x.status),('生产部','报告','2026-04-08','可归档'))

if __name__=='__main__':unittest.main(verbosity=2)
